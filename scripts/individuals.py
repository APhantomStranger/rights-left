"""
individuals.py — "The Individual Crimes" database.

A second, separate record alongside the weekly timeline: legal matters
(criminal cases, civil suits, impeachments, contempt, ethics findings,
disbarments, formal investigations) brought against individual Republican
officials, Trump administration figures, and right-wing media figures.

Everything lives in data/Individual_Crimes_Tracker.xlsx:

  Matters    one row per legal matter, with its current status and a dated
             history. This is what the site's "The Individual Crimes" page
             shows. Edit it directly in Excel whenever you like.
  Watchlist  the people the weekly search looks for: Republican members of
             Congress, Trump-appointed and Republican-appointed judges, the
             cabinet, right-wing media and podcasters, Trump associates.
             Being on the watchlist says nothing about anyone — the site
             never lists watchlist names as entries; a search for someone
             with no matters just says so.

Weekly cycle (all automatic except your review):
  gather.py   searches the news (and CourtListener, if a token is set) for
              watchlist names next to legal words, and writes what it finds
              to an "Individuals" tab in the weekly candidates sheet.
  you         mark y on the stories that are real legal developments. Put
              an existing Matter ID to UPDATE that matter (new status, new
              date), or leave it blank to ADD a new matter.
  enrich.py   fills blank Type / Offense / Status / Summary from the article.
  ingest.py   applies the approved rows to the Matters sheet.
  build_site  publishes the Matters sheet to the site.

Accuracy rules the whole thing is built around:
  - Status words follow the legal reality: "Charged" is not "Convicted".
    The site words entries the same way.
  - Matters are updated, never silently deleted: a dismissal or acquittal
    is recorded as such.
  - Only personal matters. Policy lawsuits against an official in their
    official capacity don't belong here.
"""

import datetime as dt
import html
import json
import os
import re
import sys
import time
import urllib.parse

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

TRACKER = os.path.join("data", "Individual_Crimes_Tracker.xlsx")
TAB = "Individuals"          # tab name in the weekly candidates sheet

GROUPS = ["Trump", "Cabinet & Administration", "Senate", "House",
          "Governor & State Official", "Party Official", "Judiciary",
          "Right-Wing Media", "Election Subversion & Jan. 6", "Trump Associate"]
TYPES = ["Criminal", "Civil", "Impeachment", "Contempt",
         "Ethics & Professional", "Investigation"]
OFFENSES = ["Fraud", "Election interference", "Defamation", "Obstruction of justice",
            "Contempt of Congress", "Sexual misconduct", "Campaign finance", "Conspiracy",
            "Seditious conspiracy", "Assault & violence", "Classified documents",
            "Tax crimes", "Securities & insider trading", "Bribery & corruption",
            "False statements & perjury", "Abuse of power", "Hatch Act",
            "Ethics violations", "Theft & embezzlement", "Weapons", "Drugs & DUI",
            "Foreign agent / lobbying", "Professional misconduct",
            "Civil rights violations", "Other"]
STATUSES = ["Under investigation", "Charged", "Pleaded guilty", "Convicted", "Sentenced",
            "Acquitted", "Dismissed", "Overturned", "Pardoned", "Commuted",
            "Lawsuit filed", "Judgment against", "Settled", "Articles filed",
            "Impeached", "Impeached & acquitted", "Held in contempt", "Violation found",
            "Disbarred", "License suspended", "Closed — no action", "Pending"]

# Matters sheet columns: (key, header, width)
MATTER_COLS = [
    ("id", "ID", 8), ("name", "Name", 22), ("group", "Group", 20),
    ("role", "Role", 28), ("state", "State", 6), ("matter", "Matter", 34),
    ("matter_type", "Type", 14), ("offense", "Offense", 20), ("status", "Status", 16),
    ("filed_date", "Filed", 11), ("latest_date", "Latest", 11),
    ("jurisdiction", "Jurisdiction", 28), ("case_number", "Case No.", 18),
    ("summary", "Summary", 60), ("history", "History", 60),
    ("court_url", "Court Record", 36), ("source_url", "Source", 36),
    ("source_outlet", "Outlet", 16), ("updated", "Updated", 11),
]
WATCH_COLS = [("name", "Name", 26), ("group", "Group", 24), ("role", "Role", 40),
              ("state", "State", 6), ("aliases", "Aliases (comma-separated)", 30),
              ("origin", "Added from", 16)]

# Weekly review tab columns
REVIEW_COLS = [
    ("include", "Include?", 9), ("name", "Name", 22), ("headline", "Headline", 60),
    ("outlet", "Outlet", 16), ("date", "Date", 12), ("url", "Link", 36),
    ("matter_id", "Matter ID (to update)", 12), ("open_matters", "This person's matters", 40),
    ("matter", "Matter (new)", 30), ("matter_type", "Type", 14), ("offense", "Offense", 20),
    ("status", "Status", 16), ("summary", "Summary", 50), ("court_url", "Court Record", 30),
    ("found_by", "Found by", 14),
]

NAVY, LIGHT_GREEN = "1F3864", "C6E8C6"
HDR_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", start_color=NAVY)
BODY = Font(name="Arial", size=10)
LINK_FONT = Font(name="Arial", size=10, color="0563C1", underline="single")
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))
TRUTHY = {"y", "yes", "true", "1", "x"}

# Words that make a headline about a watchlist name worth a look
LEGAL_RE = re.compile(
    r"\b(indict\w*|charg(?:e|ed|es|ing)|arrest\w*|plead\w*|pleaded|guilty|convict\w*|"
    r"sentenc\w*|acquit\w*|lawsuit\w*|sued|sues|suing|defamation|libel|impeach\w*|"
    r"contempt|subpoena\w*|ethics|investigat\w*|probe|disbar\w*|grand jury|verdict|"
    r"judgment|settle\w*|pardon\w*|commut\w*|misconduct|hatch act|inspector general|"
    r"criminal|felony|misdemeanor|prosecut\w*|trial|court|judge|fraud|bribe\w*|"
    r"censure\w*|referral|complaint|sanction\w*|warrant|raid\w*|testif\w*)\b", re.I)

# Strong words weigh more when ranking
STRONG_RE = re.compile(
    r"\b(indict\w*|charged|arrest\w*|pleaded|guilty|convict\w*|sentenc\w*|"
    r"impeach\w*|contempt|disbar\w*|verdict|grand jury|sued|lawsuit|defamation|"
    r"censure\w*|warrant)\b", re.I)

COMMON_SURNAMES = {
    "smith", "johnson", "williams", "brown", "jones", "miller", "davis", "wilson",
    "anderson", "taylor", "thomas", "moore", "martin", "jackson", "thompson", "white",
    "harris", "clark", "lewis", "robinson", "walker", "young", "allen", "king", "wright",
    "scott", "green", "baker", "adams", "nelson", "hill", "campbell", "mitchell",
    "roberts", "carter", "phillips", "evans", "turner", "torres", "parker", "collins",
    "edwards", "stewart", "morris", "murphy", "cook", "rogers", "morgan", "cooper",
    "peterson", "reed", "bailey", "bell", "kelly", "howard", "ward", "cox", "richardson",
    "wood", "watson", "brooks", "bennett", "gray", "james", "hughes", "price", "sanders",
    "myers", "long", "ross", "foster", "powell", "jenkins", "perry", "russell", "sullivan",
    "fisher", "henderson", "coleman", "simmons", "patterson", "jordan", "reynolds",
    "hamilton", "graham", "kim", "gonzales", "alexander", "ramirez", "wallace", "griffin",
    "west", "cole", "hayes", "bryant", "herrera", "gibson", "ellis", "tran", "medina",
    "wells", "webb", "warren", "burns", "lee", "lane", "hunt", "black", "stone", "mills",
    "rice", "fox", "rose", "barr", "cruz", "paul", "rubio", "johnson", "banks", "massie",
    "mace", "rose", "hunter", "loomer", "crane", "carter", "steel", "wagner", "kean",
    "sessions", "trump", "vance", "biden", "kennedy", "bush",
    # ordinary words, or surnames shared with well-known people not on the list
    "justice", "branch", "counts", "strong", "divine", "flowers", "cannon", "joseph",
    "franklin", "newsom", "stevens", "porter", "palmer", "graves", "curtis", "hudson",
    "norman", "marshall", "maxwell", "goldman", "connolly", "clarke", "schwartz", "larsen",
    "brennan", "carson", "cassidy", "kenney", "fitzgerald", "schmidt", "morales", "moreno",
    "salazar", "nichols", "erickson", "mercer", "sinatra", "shepherd", "webster", "jennings",
    "fuller", "goodwin", "haines", "barber", "berger", "dawson", "dunlap", "tiffany", "teeter",
    "rushing", "cotton", "crawford", "davidson", "downing", "duncan", "simpson", "johnston",
    "griffith", "hanlon", "higgins", "newman", "norris", "pittman", "rowland", "sweeney",
    "morrison", "buchanan", "chamberlin", "freeman", "forrest", "baxter", "benton",
    "armstrong", "alford", "mcguire", "mcfarland", "mcdowell", "mcclain", "mcmahon",
    "meredith", "marston", "rounds", "sheehy", "ludwig", "hendrix", "valderrama", "gimenez",
    "dooley", "hinson", "holcomb", "fowlkes", "doughty", "younge", "seeger", "reschenthaler",
}


# ---------------------------------------------------------------- helpers

def _s(v):
    if v is None:
        return ""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def _write_sheet(ws, cols, rows, link_keys=()):
    for c, (key, label, width) in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=c, value=label)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}1"
    for r, row in enumerate(rows, start=2):
        for c, (key, _, _) in enumerate(cols, start=1):
            v = row.get(key)
            if isinstance(v, list):
                v = "\n".join(f"{h.get('date', '')} — {h.get('event', '')}" for h in v)
            cell = ws.cell(row=r, column=c, value=v if v not in ("", None) else None)
            cell.font = BODY
            cell.alignment = WRAP
            cell.border = THIN
            if key in link_keys and v:
                cell.hyperlink = v
                cell.font = LINK_FONT


def _read_sheet(ws, cols):
    by_label = {label.lower(): key for key, label, _ in cols}
    idx = {}
    for c in range(1, ws.max_column + 1):
        h = _s(ws.cell(row=1, column=c).value).lower()
        if h in by_label:
            idx[by_label[h]] = c
    rows = []
    for r in range(2, ws.max_row + 1):
        row = {}
        for key, c in idx.items():
            cell = ws.cell(row=r, column=c)
            v = _s(cell.value)
            if not v and cell.hyperlink is not None and cell.hyperlink.target:
                v = cell.hyperlink.target
            row[key] = v
        if any(row.values()):
            row["_row"] = r
            rows.append(row)
    return rows


def parse_history(text):
    out = []
    for line in (text or "").splitlines():
        line = line.strip()
        if not line:
            continue
        m = re.match(r"^(\d{4}-\d{2}-\d{2})\s*[—–-]+\s*(.*)$", line)
        out.append({"date": m.group(1), "event": m.group(2)} if m
                   else {"date": "", "event": line})
    return out


# ---------------------------------------------------------------- tracker I/O

def load_tracker(path=TRACKER):
    """(matters, watchlist) as lists of dicts. Missing file → two empty lists."""
    if not os.path.exists(path):
        return [], []
    wb = load_workbook(path)
    matters = _read_sheet(wb["Matters"], MATTER_COLS) if "Matters" in wb.sheetnames else []
    watch = _read_sheet(wb["Watchlist"], WATCH_COLS) if "Watchlist" in wb.sheetnames else []
    for m in matters:
        m["history"] = parse_history(m.get("history"))
    return matters, watch


def save_tracker(matters, watch, path=TRACKER):
    wb = Workbook()
    ws = wb.active
    ws.title = "Matters"
    ordered = sorted(matters, key=lambda m: (m.get("latest_date") or "", m.get("id") or ""),
                     reverse=True)
    _write_sheet(ws, MATTER_COLS, ordered, link_keys=("court_url", "source_url"))
    ws2 = wb.create_sheet("Watchlist")
    _write_sheet(ws2, WATCH_COLS, sorted(watch, key=lambda w: (w.get("group", ""), w.get("name", ""))))
    about = wb.create_sheet("About")
    about.column_dimensions["A"].width = 120
    lines = [
        "The Individual Crimes — how this workbook works",
        "",
        "Matters: one row per legal matter. Status is the CURRENT status; History is one dated line per development",
        "  (format: 2026-10-12 — Indicted on two counts). The site shows this sheet, newest Latest date first.",
        "  You can edit any cell by hand. Keep IDs unique (I-001, I-002, …). Add a row to add a matter.",
        "Watchlist: the names the weekly search looks for. Being on it says nothing about anyone; the site",
        "  only lists people who have rows on Matters. Add aliases (e.g. 'RFK Jr.') to help matching.",
        "",
        "Type: " + ", ".join(TYPES),
        "Offense: " + ", ".join(OFFENSES),
        "Status: " + ", ".join(STATUSES),
    ]
    for i, t in enumerate(lines, start=1):
        c = about.cell(row=i, column=1, value=t)
        c.font = Font(name="Arial", size=11 if i == 1 else 10, bold=(i == 1))
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb.save(path)


def next_id(matters):
    nums = [int(m["id"][2:]) for m in matters if re.match(r"^I-\d+$", m.get("id") or "")]
    return f"I-{(max(nums) + 1) if nums else 1:03d}"


# ---------------------------------------------------------------- name matching

def _norm(s):
    s = s.lower().replace("’", "'")
    return re.sub(r"[^a-z0-9' ]+", " ", s).strip()


class NameMatcher:
    """Finds watchlist names in headlines. Matches the full name or an alias;
    also a distinctive surname alone ("Hegseth", "Giuliani") when no one else
    on the list shares it and it isn't a common surname."""

    def __init__(self, watch):
        self.people = {}
        surname_count = {}
        for w in watch:
            name = w.get("name") or ""
            if not name:
                continue
            last = _norm(name).split()[-1] if _norm(name).split() else ""
            if last in ("jr", "sr", "ii", "iii", "iv") and len(_norm(name).split()) > 1:
                last = _norm(name).split()[-2]
            surname_count[last] = surname_count.get(last, 0) + 1
            self.people[name] = {"row": w, "last": last,
                                 "forms": {_norm(name)} | {_norm(a) for a in
                                                           (w.get("aliases") or "").split(",") if a.strip()}}
        pats = []
        for name, p in self.people.items():
            forms = set(p["forms"])
            first_last = _norm(name).split()
            if len(first_last) >= 3:            # "Robert F. Kennedy Jr." → "robert kennedy"
                forms.add(f"{first_last[0]} {p['last']}")
            if (surname_count.get(p["last"]) == 1 and len(p["last"]) >= 6
                    and p["last"] not in COMMON_SURNAMES
                    and p["row"].get("group") != "Judiciary"):   # judges: full name only
                forms.add(p["last"])
            p["forms"] = {f for f in forms if f}
            for f in p["forms"]:
                pats.append((f, name))
        pats.sort(key=lambda t: -len(t[0]))
        self.pats = [(re.compile(r"(?<![a-z0-9])" + re.escape(f) + r"(?![a-z0-9])"), n) for f, n in pats]

    def find(self, text):
        t = _norm(text)
        hits = []
        for rx, name in self.pats:
            if name not in hits and rx.search(t):
                hits.append(name)
        return hits


# ---------------------------------------------------------------- weekly search

GNEWS = "https://news.google.com/rss/search?hl=en-US&gl=US&ceid=US:en&q="
LEGAL_TERMS = ("(indicted OR charged OR arrested OR lawsuit OR sued OR impeachment OR "
               "contempt OR convicted OR sentenced OR ethics OR investigation OR pleaded OR "
               "subpoena OR disbarred OR verdict OR defamation)")
# Catch-alls for people not on the watchlist yet
GENERAL_QUERIES = [
    'Republican (indicted OR charged OR arrested OR convicted OR sentenced)',
    'GOP (lawmaker OR congressman OR senator OR legislator) (indicted OR charged OR arrested)',
    '"Republican Party" chair (charged OR indicted OR arrested)',
    '"House Ethics Committee" Republican',
    '"articles of impeachment" (secretary OR director OR attorney general OR judge)',
    '"held in contempt" Trump administration official',
    'conservative (podcaster OR commentator OR influencer OR host) (charged OR arrested OR sued OR indicted)',
    '"Fox News" host (sued OR lawsuit OR charged)',
    'Newsmax (sued OR lawsuit OR charged OR settlement)',
    '"Jan. 6" pardoned (arrested OR charged)',
]


def _http_get(url, timeout=25, headers=None):
    import requests
    h = {"User-Agent": "Mozilla/5.0 (RightsLeft tracker; +https://rightsleft.org)"}
    h.update(headers or {})
    return requests.get(url, headers=h, timeout=timeout)


def _parse_rss(xml):
    items = []
    for it in re.findall(r"<item>(.*?)</item>", xml, re.S):
        def tag(t):
            m = re.search(rf"<{t}[^>]*>(.*?)</{t}>", it, re.S)
            return html.unescape(re.sub(r"<!\[CDATA\[|\]\]>", "", m.group(1))).strip() if m else ""
        title = tag("title")
        src = re.search(r'<source url="([^"]*)">(.*?)</source>', it, re.S)
        outlet = html.unescape(src.group(2)).strip() if src else ""
        if outlet and title.endswith(" - " + outlet):
            title = title[: -len(" - " + outlet)]
        pub = tag("pubDate")
        try:
            d = dt.datetime.strptime(pub[:25], "%a, %d %b %Y %H:%M:%S").date()
        except ValueError:
            d = None
        items.append({"headline": title, "url": tag("link"), "outlet": outlet, "date": d})
    return items


def _gnews(query, days, log):
    url = GNEWS + urllib.parse.quote(f"{query} when:{days}d")
    for attempt in range(3):
        try:
            r = _http_get(url)
            if r.status_code == 200:
                return _parse_rss(r.text)
            if r.status_code in (429, 503):
                time.sleep(10 * (attempt + 1))
                continue
            log(f"  ! news search HTTP {r.status_code}")
            return []
        except Exception as ex:
            log(f"  ! news search failed ({ex.__class__.__name__})")
            time.sleep(3)
    return []


def _decode_google_link(url):
    if "news.google.com" not in url:
        return url
    try:
        from googlenewsdecoder import gnewsdecoder
        r = gnewsdecoder(url, interval=0.5)
        if r.get("success"):
            return r["decoded_url"]
    except Exception:
        pass
    return url


def _courtlistener(watch, since, token, log):
    """New federal court dockets naming watchlist people (RECAP archive).
    Needs a free CourtListener API token in COURTLISTENER_TOKEN."""
    out = []
    for w in watch:
        name = w.get("name")
        if not name:
            continue
        q = urllib.parse.urlencode({"type": "r", "q": f'"{name}"',
                                    "filed_after": since.strftime("%m/%d/%Y"),
                                    "order_by": "dateFiled desc"})
        try:
            r = _http_get("https://www.courtlistener.com/api/rest/v4/search/?" + q,
                          headers={"Authorization": f"Token {token}"})
            if r.status_code != 200:
                if r.status_code in (401, 403):
                    log("  ! CourtListener token rejected — skipping court search")
                    return out
                continue
            for res in (r.json() or {}).get("results", [])[:5]:
                case = res.get("caseName") or ""
                last = _norm(name).split()[-1]
                if last not in _norm(case):          # name must be a party, not just mentioned
                    continue
                link = "https://www.courtlistener.com" + (res.get("docket_absolute_url") or "")
                out.append({"name": name, "headline": f"New federal case: {case}",
                            "outlet": f"CourtListener · {res.get('court') or ''}".strip(" ·"),
                            "date": (res.get("dateFiled") or "")[:10], "url": link,
                            "court_url": link, "found_by": "Court records", "_score": 6})
        except Exception:
            continue
        time.sleep(0.2)
    return out


def weekly_search(watch, matters, days=8, log=print, max_items=250, decode=True,
                  max_per_person=4, max_general=60):
    """Candidate rows for the weekly Individuals tab, strongest first."""
    matcher = NameMatcher(watch)
    names = [w["name"] for w in watch if w.get("name")]
    seen, rows = set(), []
    oldest = dt.date.today() - dt.timedelta(days=days + 2)

    def add(it, found_by, general=False):
        url = it.get("url") or ""
        if it.get("date") and it["date"] < oldest:      # search engines' date filters leak
            return
        key = (it.get("headline") or "").lower()[:120]
        if not url or key in seen:
            return
        text = it.get("headline") or ""
        if not LEGAL_RE.search(text):
            return
        hits = matcher.find(text)
        if not hits and not general:
            return
        seen.add(key)
        score = 2 + 3 * len(STRONG_RE.findall(text)) + (2 if hits else 0)
        rows.append({"name": ", ".join(hits), "headline": text, "outlet": it.get("outlet", ""),
                     "date": it["date"].isoformat() if it.get("date") else "",
                     "url": url, "found_by": found_by if hits else "General search",
                     "_score": score})

    # 1) watchlist names, six at a time, next to legal words
    batch = 6
    for i in range(0, len(names), batch):
        chunk = names[i:i + batch]
        q = "(" + " OR ".join(f'"{n}"' for n in chunk) + ") " + LEGAL_TERMS
        for it in _gnews(q, days, log):
            add(it, "Name search")
        time.sleep(1.2)
    log(f"  Individuals: {len(rows)} name match(es) from {len(names)} watchlist names")

    # 2) catch-alls for people not on the list yet
    for q in GENERAL_QUERIES:
        for it in _gnews(q, days, log):
            add(it, "General search", general=True)
        time.sleep(1.2)

    # 3) optional: new federal dockets
    token = os.environ.get("COURTLISTENER_TOKEN")
    if token:
        since = dt.date.today() - dt.timedelta(days=days)
        cl = _courtlistener(watch, since, token, log)
        for c in cl:
            if c["url"] not in {r["url"] for r in rows}:
                rows.append(c)
        log(f"  Individuals: {len(cl)} new federal docket(s) from CourtListener")
    else:
        log("  Individuals: court-record search OFF (no COURTLISTENER_TOKEN)")

    rows.sort(key=lambda r: r.get("date") or "", reverse=True)    # newest first…
    rows.sort(key=lambda r: -r["_score"])                         # …within each score
    # One big story can produce a dozen near-identical headlines; keep the
    # strongest few per person so the tab stays reviewable.
    per, capped, general = {}, [], 0
    for r in rows:
        if not r["name"]:                 # catch-all results: cap so they don't crowd out names
            general += 1
            if general <= max_general:
                capped.append(r)
            continue
        per[r["name"]] = per.get(r["name"], 0) + 1
        if per[r["name"]] <= max_per_person:
            capped.append(r)
    rows = capped[:max_items]

    # Fill in each person's existing matters so updating one is a copy-paste
    by_name = {}
    for m in matters:
        by_name.setdefault(m.get("name"), []).append(m)
    for r in rows:
        ms = [m for n in r["name"].split(", ") for m in by_name.get(n, [])]
        r["open_matters"] = "\n".join(f"{m['id']}: {m.get('matter', '')} ({m.get('status', '')})"
                                      for m in ms)
        if decode and "news.google.com" in r["url"]:
            r["url"] = _decode_google_link(r["url"])
        r.pop("_score", None)
    return rows


# ---------------------------------------------------------------- weekly review tab

def write_review_tab(wb, rows):
    """Add the Individuals tab to a weekly candidates workbook. Dropdowns
    for Include?, Type, Offense and Status point at the hidden Lists sheet
    (columns C, D, E) which this adds to."""
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.worksheet.datavalidation import DataValidation

    lists = wb["Lists"] if "Lists" in wb.sheetnames else wb.create_sheet("Lists")
    lists.sheet_state = "hidden"
    for col, values in ((3, TYPES), (4, OFFENSES), (5, STATUSES)):
        for i, v in enumerate(values, start=1):
            lists.cell(row=i, column=col, value=v)

    ws = wb.create_sheet(TAB)
    _write_sheet(ws, REVIEW_COLS, rows, link_keys=("url", "court_url"))
    ws.freeze_panes = "C2"
    last = max(len(rows) + 1, 2) + 40
    keys = [k for k, _, _ in REVIEW_COLS]
    col = lambda k: get_column_letter(keys.index(k) + 1)

    dv = DataValidation(type="list", formula1='"y,n"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"A2:A{last}")
    for key, lcol, n in (("matter_type", "C", len(TYPES)), ("offense", "D", len(OFFENSES)),
                         ("status", "E", len(STATUSES))):
        d = DataValidation(type="list", formula1=f"Lists!${lcol}$1:${lcol}${n}", allow_blank=True)
        ws.add_data_validation(d)
        d.add(f"{col(key)}2:{col(key)}{last}")
    ws.conditional_formatting.add(
        f"A2:{get_column_letter(len(REVIEW_COLS))}{last}",
        FormulaRule(formula=['LOWER($A2)="y"'], fill=PatternFill("solid", start_color=LIGHT_GREEN)))
    return ws


def read_review_rows(path):
    if not path or not path.lower().endswith(".xlsx") or not os.path.exists(path):
        return []
    wb = load_workbook(path, data_only=True)
    if TAB not in wb.sheetnames:
        return []
    return _read_sheet(wb[TAB], REVIEW_COLS)


# ---------------------------------------------------------------- ingest

def _pick(value, allowed):
    if not value:
        return ""
    for a in allowed:
        if a.lower() == value.strip().lower():
            return a
    return value.strip()


def apply_review(cand_path, tracker=TRACKER, today=None, log=print):
    """Apply the approved Individuals rows from a weekly sheet to the
    tracker. A row with a Matter ID updates that matter; a row without one
    adds a new matter. Returns (added, updated)."""
    rows = [r for r in read_review_rows(cand_path)
            if (r.get("include") or "").strip().lower() in TRUTHY]
    if not rows:
        return 0, 0
    today = today or dt.date.today()
    matters, watch = load_tracker(tracker)
    by_id = {m["id"]: m for m in matters}
    watch_by_name = {w["name"]: w for w in watch}
    added = updated = 0
    for r in rows:
        date = r.get("date") or today.isoformat()
        try:
            date = dt.date.fromisoformat(date[:10]).isoformat()
        except ValueError:
            date = today.isoformat()
        status = _pick(r.get("status"), STATUSES)
        event = (status + ": " if status else "") + (r.get("headline") or r.get("summary") or "update")
        mid = (r.get("matter_id") or "").strip().upper()
        if mid and mid in by_id:
            m = by_id[mid]
            if status:
                m["status"] = status
            if date >= (m.get("latest_date") or ""):
                m["latest_date"] = date
            m["history"] = sorted(m.get("history", []) + [{"date": date, "event": event}],
                                  key=lambda h: h.get("date") or "")
            for k in ("court_url",):
                if r.get(k):
                    m[k] = r[k]
            if r.get("url"):
                m["source_url"], m["source_outlet"] = r["url"], r.get("outlet", "")
            if r.get("summary"):
                m["summary"] = r["summary"]
            m["updated"] = today.isoformat()
            updated += 1
            log(f"  · updated {mid} ({m.get('name')}): {status or 'new development'}")
            continue
        if mid:
            log(f"  ! Matter ID {mid} not found — adding as a new matter instead")
        name = (r.get("name") or "").split(",")[0].strip()
        if not name:
            log(f"  ! skipped (no name): {(r.get('headline') or '')[:70]}")
            continue
        w = watch_by_name.get(name, {})
        m = {"id": next_id(matters), "name": name, "group": w.get("group", ""),
             "role": w.get("role", ""), "state": w.get("state", ""),
             "matter": r.get("matter") or (r.get("headline") or "")[:80],
             "matter_type": _pick(r.get("matter_type"), TYPES),
             "offense": _pick(r.get("offense"), OFFENSES),
             "status": status, "filed_date": date, "latest_date": date,
             "jurisdiction": "", "case_number": "", "summary": r.get("summary", ""),
             "history": [{"date": date, "event": event}],
             "court_url": r.get("court_url", ""), "source_url": r.get("url", ""),
             "source_outlet": r.get("outlet", ""), "updated": today.isoformat()}
        matters.append(m)
        by_id[m["id"]] = m
        if name not in watch_by_name:
            nw = {"name": name, "group": "", "role": "", "state": "", "aliases": "",
                  "origin": "weekly review"}
            watch.append(nw)
            watch_by_name[name] = nw
        added += 1
        log(f"  · added {m['id']} ({name}): {m['matter'][:60]}")
    save_tracker(matters, watch, tracker)
    return added, updated


# ---------------------------------------------------------------- site payload

def site_payload(tracker=TRACKER):
    matters, watch = load_tracker(tracker)
    matters = [m for m in matters if m.get("name")]
    for m in matters:
        m.pop("_row", None)
    matters.sort(key=lambda m: (m.get("latest_date") or "", m.get("id") or ""), reverse=True)
    counts = {}
    for m in matters:
        if m.get("offense"):
            counts[m["offense"]] = counts.get(m["offense"], 0) + 1
    top = sorted(counts.items(), key=lambda kv: (-kv[1], kv[0]))[:10]
    people = sorted({m["name"] for m in matters})
    status_counts = {}
    for m in matters:
        status_counts[m.get("status") or "—"] = status_counts.get(m.get("status") or "—", 0) + 1
    watch_names = sorted({w["name"] for w in watch if w.get("name")} - set(people))
    updated = max((m.get("updated") or m.get("latest_date") or "" for m in matters), default="")
    return {
        "matters": [{k: m.get(k, "") for k in
                     ("id", "name", "group", "role", "state", "matter", "matter_type", "offense",
                      "status", "filed_date", "latest_date", "jurisdiction", "case_number",
                      "summary", "history", "court_url", "source_url", "source_outlet")}
                    for m in matters],
        "offenses": [{"category": k, "count": v} for k, v in top],
        "people": len(people),
        "statuses": status_counts,
        "watch": watch_names,
        "watch_total": len(watch),
        "updated": updated,
        "groups": [g for g in GROUPS if any(m.get("group") == g for m in matters)],
    }


def recent_changes(tracker=TRACKER, days=7, today=None):
    """Matters added or updated in the last `days` — for the newsletter."""
    today = today or dt.date.today()
    cutoff = (today - dt.timedelta(days=days)).isoformat()
    matters, _ = load_tracker(tracker)
    out = [m for m in matters if (m.get("updated") or "") >= cutoff]
    out.sort(key=lambda m: m.get("latest_date") or "", reverse=True)
    return out
