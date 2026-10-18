"""
note.py — your personal note for the weekly newsletter.

Every weekly candidates workbook gets a "My Note" tab. Whatever you type
there rides along with the approved stories:

  gather.py   creates the tab (and keeps what you typed if gather re-runs)
  ingest.py   copies it to data/notes/<week Monday>.md
  newsletter.py  puts it at the top of that week's email, above the list

Leave the tab blank and the newsletter goes out exactly as before.

You can also skip the spreadsheet and upload a Markdown file straight to
data/notes/ on GitHub, named for the week's Monday (e.g. 2026-10-12.md),
before running Ingest. Format:

    # Optional title (first line, starting with "# ")

    First paragraph.

    Second paragraph. Bare links like https://example.com become clickable.

    — Chris

A note typed in the sheet replaces an uploaded file for the same week.
"""

import html, os, re

from openpyxl.styles import Alignment, Font, PatternFill

NOTE_TAB = "My Note"
NOTES_DIR = os.path.join("data", "notes")
DEFAULT_SIGNOFF = "— Chris"

# Fixed cells on the tab (column B holds what you type)
TITLE_CELL = "B4"
SIGNOFF_CELL = "B5"
FIRST_PARA_ROW = 8
PARA_ROWS = 40          # rows that get the input shading

_INPUT_FILL = PatternFill("solid", start_color="FFF8DC")


def write_note_tab(wb, note=None):
    """Add the My Note tab to a workbook being built. `note` is a dict from
    read_note() — passed in when gather re-runs over a sheet you've already
    started, so nothing you typed is lost."""
    note = note or {}
    ws = wb.create_sheet(NOTE_TAB)
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 110
    bold = Font(name="Arial", size=12, bold=True)
    body = Font(name="Arial", size=10)
    hint = Font(name="Arial", size=10, italic=True, color="666666")
    wrap = Alignment(wrap_text=True, vertical="top")

    ws["A1"] = "Your note for this week's newsletter (optional)"
    ws["A1"].font = bold
    ws["A2"] = ("Type a title and your note in the yellow cells, one paragraph per row. "
                "It goes at the top of this week's email, above the list of stories. "
                "Leave it all blank to send the newsletter without a note.")
    ws["A2"].font = hint

    ws["A4"] = "Title:"
    ws["A5"] = "Sign-off:"
    ws["A7"] = "Your note:"
    for ref in ("A4", "A5", "A7"):
        ws[ref].font = Font(name="Arial", size=10, bold=True)
    ws["B7"] = "One paragraph per row, starting on the next row. Blank rows are skipped."
    ws["B7"].font = hint

    ws[TITLE_CELL] = note.get("title") or None
    ws[SIGNOFF_CELL] = note.get("signoff") or DEFAULT_SIGNOFF
    for ref in (TITLE_CELL, SIGNOFF_CELL):
        ws[ref].font = body
        ws[ref].fill = _INPUT_FILL

    paras = list(note.get("paragraphs") or [])
    for i in range(max(PARA_ROWS, len(paras))):
        c = ws.cell(row=FIRST_PARA_ROW + i, column=2,
                    value=paras[i] if i < len(paras) else None)
        c.font = body
        c.alignment = wrap
        c.fill = _INPUT_FILL
    return ws


def read_note(path):
    """The note typed on a candidates workbook's My Note tab, as
    {"title", "signoff", "paragraphs"}; None if there's no tab or no text."""
    if not path or not path.lower().endswith(".xlsx") or not os.path.exists(path):
        return None
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    if NOTE_TAB not in wb.sheetnames:
        return None
    ws = wb[NOTE_TAB]
    clean = lambda v: "" if v is None else str(v).strip()
    paras = []
    for r in range(FIRST_PARA_ROW, ws.max_row + 1):
        v = clean(ws.cell(row=r, column=2).value)
        if v:
            # a paragraph typed with Alt+Enter line breaks splits on blank lines
            paras.extend(p.strip() for p in re.split(r"\n\s*\n", v) if p.strip())
    note = {"title": clean(ws[TITLE_CELL].value),
            "signoff": clean(ws[SIGNOFF_CELL].value),
            "paragraphs": paras}
    return note if (note["title"] or paras) else None


def to_markdown(note):
    parts = []
    if note.get("title"):
        parts.append("# " + note["title"])
    parts.extend(note.get("paragraphs") or [])
    if note.get("signoff") and (note.get("paragraphs") or note.get("title")):
        parts.append(note["signoff"])
    return "\n\n".join(parts) + "\n"


def from_markdown(text):
    blocks = [b.strip() for b in re.split(r"\n\s*\n", text or "") if b.strip()]
    title = ""
    if blocks and blocks[0].startswith("# "):
        title = blocks.pop(0)[2:].strip()
    signoff = ""
    if blocks and re.match(r"^(—|–|-{1,2})\s*\S", blocks[-1]) and len(blocks[-1]) < 80:
        signoff = blocks.pop()
    return {"title": title, "signoff": signoff, "paragraphs": blocks}


def note_path(monday, notes_dir=NOTES_DIR):
    return os.path.join(notes_dir, f"{monday:%Y-%m-%d}.md")


def save_note(note, monday, notes_dir=NOTES_DIR):
    os.makedirs(notes_dir, exist_ok=True)
    path = note_path(monday, notes_dir)
    with open(path, "w", encoding="utf-8") as f:
        f.write(to_markdown(note))
    return path


def load_note(monday, notes_dir=NOTES_DIR):
    path = note_path(monday, notes_dir)
    if not os.path.exists(path):
        return None
    with open(path, encoding="utf-8") as f:
        note = from_markdown(f.read())
    return note if (note["title"] or note["paragraphs"]) else None


_URL_RE = re.compile(r"(https?://[^\s<]+[^\s<.,;:!?)\]'\"])")


def paragraph_html(text, link_color):
    """Escape a paragraph for email, keep line breaks, make bare URLs links."""
    out = html.escape(text)
    out = _URL_RE.sub(
        lambda m: f'<a href="{m.group(1)}" style="color:{link_color};">{m.group(1)}</a>', out)
    return out.replace("\n", "<br>")
